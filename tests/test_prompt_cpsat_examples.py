"""The cpsat_python_solution_workflow prompt's code examples must be copyable as-is.

These tests validate the examples AFTER template rendering (when doubled
braces have become normal Python braces) and run them in-process against the
pinned OR-Tools dependency — no subprocess, no network, no managed runtime.
"""

from __future__ import annotations

import contextlib
import dataclasses
import io
import json
import os
import re
import sys
import textwrap
import zipfile
from collections.abc import Generator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from openconstraint_mcp.protocol_text.prompts import SOLVE_CPSAT_PYTHON_PROMPT

_CONTRACT_KEYS = {"status", "objective", "solution", "best_objective_bound"}

_SOLVE_SIGNATURE = "def solve(instance: ProblemInstance) -> Solution:"

_REPLAY_ENV_VARS = ("OPENCONSTRAINT_MCP_CPSAT_SEED", "OPENCONSTRAINT_MCP_CPSAT_CONFIG")


@contextlib.contextmanager
def _isolated_replay_env() -> Generator[None]:
    # run_cpsat_python clears both replay-protocol env vars before executing a
    # script; mirror that here so an inherited OPENCONSTRAINT_MCP_CPSAT_SEED
    # (e.g. a non-integer value) or OPENCONSTRAINT_MCP_CPSAT_CONFIG (e.g. a
    # path to a missing file) cannot change what an example does.
    saved = {name: os.environ.pop(name, None) for name in _REPLAY_ENV_VARS}
    try:
        yield
    finally:
        for name, value in saved.items():
            if value is not None:
                os.environ[name] = value


def _rendered_code_fences() -> list[str]:
    rendered = SOLVE_CPSAT_PYTHON_PROMPT.format(problem="toy problem")
    fences: list[str] = []
    inside = False
    buf: list[str] = []
    for line in rendered.splitlines():
        # startswith, not ==, so an opening fence with a language tag
        # (```python) still toggles instead of silently joining the body.
        if line.strip().startswith("```"):
            if inside:
                fences.append(textwrap.dedent("\n".join(buf)))
                buf = []
            inside = not inside
        elif inside:
            buf.append(line)
    return fences


def _run_capturing_stdout(source: str) -> list[str]:
    out = io.StringIO()
    with _isolated_replay_env(), contextlib.redirect_stdout(out):
        # The example guards its main() call on __name__, which an exec
        # against bare globals resolves through builtins to "builtins" —
        # so without this injection nothing runs and stdout stays empty.
        exec(compile(source, "<prompt-example>", "exec"), {"__name__": "__main__"})
    return [line for line in out.getvalue().strip().splitlines() if line]


def _define_example(source: str) -> dict[str, Any]:
    """Exec a fence WITHOUT ``__name__``, so it defines names and runs no main()."""
    namespace: dict[str, Any] = {}
    exec(compile(source, "<prompt-example>", "exec"), namespace)
    return namespace


def _infeasible_record() -> tuple[dict[str, Any], Any]:
    """Drive the main example's solve() infeasible and return its Solution record.

    OR-Tools returns 0.0 (not an exception) from objective_value and
    best_objective_bound on an infeasible solve, so the example's status guards
    — not the properties themselves — are what keep fabricated values out of
    both the record and the emitted JSON.
    """
    namespace = _define_example(_rendered_code_fences()[0])
    instance = namespace["parse_input"](namespace["read_input"]())
    # Requiring every item exceeds the capacity, so the coverage constraint
    # makes the instance genuinely infeasible.
    infeasible = dataclasses.replace(instance, min_items=len(instance.items))
    with _isolated_replay_env():
        return namespace, namespace["solve"](infeasible)


def test_cpsat_prompt_code_fences_have_no_placeholders() -> None:
    fences = _rendered_code_fences()
    assert len(fences) == 3, (
        "expected the main example, the callback variant, and the xlsx read_input"
    )
    for fence in fences:
        assert "..." not in fence, "a copyable example must not contain placeholders"


def test_cpsat_prompt_main_example_runs_and_emits_contract_json() -> None:
    main = _rendered_code_fences()[0]

    payload = json.loads(_run_capturing_stdout(main)[-1])

    assert set(payload) == _CONTRACT_KEYS
    assert payload["status"] == "optimal"
    assert payload["objective"] == 15.0
    assert payload["solution"] == {"selected": ["radio", "lamp", "clock"]}


def test_cpsat_prompt_main_example_returns_a_normalized_record_when_infeasible() -> None:
    _, record = _infeasible_record()

    assert record.status == "infeasible"
    assert record.selected is None, "absence must be None, never an empty list"
    assert record.objective is None
    assert record.best_objective_bound is None


def test_cpsat_prompt_serializer_emits_no_fabricated_values_for_an_infeasible_record() -> None:
    namespace, record = _infeasible_record()

    payload = namespace["serialize_solution"](record)

    assert set(payload) == _CONTRACT_KEYS
    assert payload["status"] == "infeasible"
    assert payload["objective"] is None
    assert payload["solution"] == {}
    assert payload["best_objective_bound"] is None


def test_cpsat_prompt_main_example_reports_no_objective_for_a_feasibility_model() -> None:
    # Both value guards also test model.has_objective(), and the infeasible and
    # unknown paths reach None through the STATUS half alone. Neutralize the
    # objective — a same-line token swap that leaves the indentation intact —
    # and a genuinely solved run must still report no objective.
    main = _rendered_code_fences()[0]
    marker = "model.maximize(total_value)"
    assert marker in main
    namespace = _define_example(main.replace(marker, "pass"))

    with _isolated_replay_env():
        record = namespace["solve"](namespace["parse_input"](namespace["read_input"]()))

    assert record.selected, "a feasibility model still reports its decision values"
    assert record.objective is None
    assert record.best_objective_bound is None


def test_cpsat_prompt_main_example_keeps_a_distinct_unknown_bound_guard() -> None:
    # UNKNOWN is not reliably producible from a knapsack this small, so pin the
    # bound guard's breadth at the source: it must stay wider than has_solution.
    main = _rendered_code_fences()[0]

    assert "bound_states = (cp_model.OPTIMAL, cp_model.FEASIBLE, cp_model.UNKNOWN)" in main
    assert "if model.has_objective() and status_code in bound_states" in main


def test_cpsat_prompt_serializer_keeps_the_bound_for_an_unknown_without_an_incumbent() -> None:
    namespace = _define_example(_rendered_code_fences()[0])
    record = namespace["Solution"](status="unknown", best_objective_bound=29.0)

    payload = namespace["serialize_solution"](record)

    assert set(payload) == _CONTRACT_KEYS
    assert payload["status"] == "unknown"
    assert payload["objective"] is None
    assert payload["solution"] == {}
    assert payload["best_objective_bound"] == 29.0


def test_cpsat_prompt_serializer_emits_a_legitimate_empty_selection_as_an_answer() -> None:
    # None means "no incumbent"; an EMPTY list is a real answer a coverage-free
    # instance can prove optimal. A truthiness check in the serializer would
    # collapse the two and make a proven optimum look like no answer at all.
    namespace = _define_example(_rendered_code_fences()[0])
    record = namespace["Solution"](status="optimal", selected=[], objective=0.0)

    payload = namespace["serialize_solution"](record)

    assert payload["solution"] == {"selected": []}


def test_cpsat_prompt_main_example_reads_the_seed_from_the_environment() -> None:
    # The prompt's prose teaches the seed protocol; the example must DO it,
    # since save_verified_cpsat_python's replay works by setting this env var.
    # num_workers is keyed on the config-driven literal, scoped to this fence
    # alone: the same literal already exists in step 6's experiment guidance,
    # but that line sits outside the fence, so scoping it here still has bite.
    main = _rendered_code_fences()[0]

    assert 'os.environ.get("OPENCONSTRAINT_MCP_CPSAT_SEED", "42")' in main
    assert 'solver.parameters.num_workers = config.get("num_workers", 1)' in main


def test_cpsat_prompt_callback_example_repeats_the_main_examples_guards() -> None:
    # The callback fence replaces the WHOLE solve(), so it carries its own copy
    # of the status guards. Nothing else compares the two copies, so a drifted
    # one — say a callback variant that lost model.has_objective() — would ship
    # self-contradictory guidance with every test still green.
    main, callback = _rendered_code_fences()[:2]
    start = "status_map = {"
    end = "def serialize_solution("
    assert start in main
    assert start in callback
    assert end in main

    main_guards = main[main.index(start) : main.index(end)]
    callback_guards = callback[callback.index(start) :]

    assert main_guards.strip() == callback_guards.strip()


def _composed_callback_script() -> str:
    """Splice the callback fence's solve() into the main fence, as the prompt says to."""
    main, callback = _rendered_code_fences()[:2]
    assert main.count(_SOLVE_SIGNATURE) == 1
    assert callback.startswith(_SOLVE_SIGNATURE)
    return (
        main[: main.index(_SOLVE_SIGNATURE)]
        + callback.strip()
        + "\n\n\n"
        + main[main.index("def serialize_solution(") :]
    )


def test_cpsat_prompt_callback_example_replaces_the_whole_solve_function() -> None:
    # The prompt instructs replacing the whole solve() function with the
    # callback variant, so the two fences must compose into one runnable script
    # that emits intermediate JSON lines before the authoritative final line.
    lines = _run_capturing_stdout(_composed_callback_script())

    assert len(lines) >= 2, "callback should emit at least one intermediate line"
    intermediate = json.loads(lines[0])
    assert intermediate["status"] == "feasible"
    assert set(intermediate) == _CONTRACT_KEYS
    final = json.loads(lines[-1])
    assert final["status"] == "optimal"


def test_cpsat_prompt_callback_bounds_intermediate_bytes_but_keeps_the_final_result() -> None:
    source = _composed_callback_script().replace(
        "self._remaining_output_bytes = 512 * 1024",
        "self._remaining_output_bytes = 1",
    )

    lines = _run_capturing_stdout(source)

    assert len(lines) == 1
    assert json.loads(lines[0])["status"] == "optimal"


def test_cpsat_prompt_callback_example_still_streams_under_a_configured_search_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Regression: the callback used to be gated on the config's CP-SAT limit, so a
    # CONFIGURED run emitted nothing until solve() returned. A child tree-killed at
    # the executor's deadline before that left no JSON on stdout at all, and timeout
    # recovery reported no incumbent — discarding every solution CP-SAT had found.
    # A search limit bounds search alone and is never validated against the executor
    # deadline, which the script cannot see, so streaming must stay unconditional.
    # Every other example test runs with the config env var cleared, which is why
    # this path went uncovered.
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps({"solver_time_limit_seconds": 30}), encoding="utf-8")
    monkeypatch.setenv("OPENCONSTRAINT_MCP_CPSAT_CONFIG", str(config_path))

    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        exec(  # noqa: S102 - the prompt's own example, executed to prove it works
            compile(_composed_callback_script(), "<prompt-example>", "exec"),
            {"__name__": "__main__"},
        )
    lines = [line for line in out.getvalue().strip().splitlines() if line]

    assert len(lines) >= 2, "a configured search limit must not suppress the stream"


# --- the xlsx read_input() example ------------------------------------------------

_XLSX_FENCE_INDEX = 2


def _workbook(path: Path, records: list[list[object]], *, dimension: str | None = None) -> Path:
    """Write ``records`` to a sheet named ``Orders``, optionally faking its dimension ref.

    openpyxl always records the true extent, so reproducing the understated ref
    that other writers emit means rewriting the stored sheet XML.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    for record in records:
        sheet.append(record)
    book.save(path)
    if dimension is None:
        return path
    target = path.with_name(f"understated_{path.name}")
    with zipfile.ZipFile(path) as src_zip, zipfile.ZipFile(target, "w") as dst_zip:
        for item in src_zip.infolist():
            data = src_zip.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data, count = re.subn(
                    rb'<dimension ref="[^"]*"\s*/>',
                    f'<dimension ref="{dimension}" />'.encode(),
                    data,
                )
                assert count == 1, "sheet XML carried no <dimension> to rewrite"
            dst_zip.writestr(item, data)
    return target


def _read_via_example(path: Path, monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
    namespace = _define_example(_rendered_code_fences()[_XLSX_FENCE_INDEX])
    monkeypatch.setattr(sys, "argv", ["model.py", str(path)])
    result: list[dict[str, Any]] = namespace["read_input"]()
    return result


def test_cpsat_prompt_xlsx_example_reads_every_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "d.xlsx", [["name", "qty"], ["widget", 3], ["gadget", 10]])

    assert _read_via_example(path, monkeypatch) == [
        {"name": "widget", "qty": 3},
        {"name": "gadget", "qty": 10},
    ]


def test_cpsat_prompt_xlsx_example_survives_an_understated_dimension_ref(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The whole reason the example carries reset_dimensions(): without it
    # openpyxl trims iteration to this ref and the sheet reads as header-only,
    # silently solving a smaller instance.
    path = _workbook(
        tmp_path / "d.xlsx",
        [["name", "qty"], ["widget", 3], ["gadget", 10]],
        dimension="A1:A1",
    )

    assert _read_via_example(path, monkeypatch) == [
        {"name": "widget", "qty": 3},
        {"name": "gadget", "qty": 10},
    ]


def test_cpsat_prompt_xlsx_example_keeps_a_dropped_column_when_the_ref_is_narrow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The trim hits the header row too, so the missing column leaves no width
    # mismatch behind for a zip(strict=True) check to notice.
    path = _workbook(
        tmp_path / "d.xlsx", [["name", "qty", "due"], ["widget", 3, "friday"]], dimension="A1:B2"
    )

    assert _read_via_example(path, monkeypatch) == [{"name": "widget", "qty": 3, "due": "friday"}]


def test_cpsat_prompt_xlsx_example_pads_a_row_whose_last_cell_is_blank(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # At natural width such a row is SHORT, so a bare strict=True would reject
    # an ordinary sheet rather than reading it.
    path = _workbook(tmp_path / "d.xlsx", [["name", "qty"], ["widget", None]])

    assert _read_via_example(path, monkeypatch) == [{"name": "widget", "qty": None}]


def test_cpsat_prompt_xlsx_example_skips_a_trailing_blank_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "d.xlsx", [["name", "qty"], ["widget", 3], [None, None]])

    assert _read_via_example(path, monkeypatch) == [{"name": "widget", "qty": 3}]


def test_cpsat_prompt_xlsx_example_rejects_duplicate_headers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # dict() would keep only the last of the two, dropping a column in silence.
    path = _workbook(tmp_path / "d.xlsx", [["name", "name"], ["widget", "gadget"]])

    with pytest.raises(ValueError, match="duplicate column names"):
        _read_via_example(path, monkeypatch)


def test_cpsat_prompt_xlsx_example_rejects_a_row_wider_than_the_header(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "d.xlsx", [["name", "qty"], ["widget", 3, "extra"]])

    with pytest.raises(ValueError, match="cells under"):
        _read_via_example(path, monkeypatch)
